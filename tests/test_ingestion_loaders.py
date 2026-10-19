from __future__ import annotations

import unicodedata

import openpyxl
import pytest
from docx import Document
from pptx import Presentation
from pypdf import PdfWriter

from ingestion.corpus import corpus_hash, load_documents
from ingestion.loaders import LoadError, load_docx, load_pdf, load_pptx, load_xlsx
from ingestion.pipeline import build_chunks, chunk_lines, ingest
from rag_core.store import open_collection
from tests.conftest import fake_embed


def test_load_docx_keeps_paragraphs_and_table_rows_in_body_order(tmp_path):
    path = tmp_path / "a.docx"
    doc = Document()
    doc.add_paragraph("before")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "x"
    table.rows[0].cells[1].text = "y"
    doc.add_paragraph("after")
    doc.save(path)

    assert load_docx(path)[0].text == "before\nx | y\nafter"


def test_load_pptx_puts_labelled_notes_after_slide_text_with_slide_location(tmp_path):
    path = tmp_path / "a.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Title"
    slide.placeholders[1].text = "Body"
    slide.notes_slide.notes_text_frame.text = "spoken"
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(path)

    sections = load_pptx(path)

    assert sections[0].text == "Title\nBody\nGhi chú:\nspoken"
    assert sections[0].location == {"slide": 1}
    assert sections[1].text == ""


def test_load_pdf_returns_one_section_per_page_and_blank_page_makes_no_chunk(tmp_path):
    path = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(200, 200)
    writer.write(path)

    sections = load_pdf(path)

    assert [s.location for s in sections] == [{"page": 1}]
    assert build_chunks(tmp_path, 40, 10) == []


def write_xlsx(path, rows, sheet="Q1"):
    workbook = openpyxl.Workbook()
    workbook.active.title = sheet
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)
    return workbook


def test_load_xlsx_renders_each_row_with_its_headers_and_skips_hidden_sheets(tmp_path):
    path = tmp_path / "a.xlsx"
    workbook = write_xlsx(path, [["name", "days"], ["An", 12], ["Binh", 5]])
    workbook.create_sheet("secret").sheet_state = "hidden"
    workbook.save(path)

    sections = load_xlsx(path)

    assert [s.location for s in sections] == [{"sheet": "Q1"}]
    assert sections[0].text == "Sheet Q1: name: An; days: 12\nSheet Q1: name: Binh; days: 5"


def test_xlsx_chunks_never_split_a_row(tmp_path):
    rows = [["name", "days"]] + [[f"person{i}", i] for i in range(10)]
    write_xlsx(tmp_path / "a.xlsx", rows)

    chunks = build_chunks(tmp_path, size=70, overlap=10)

    assert len(chunks) > 1
    assert all(line.startswith("Sheet Q1: name: ") for c in chunks for line in c.text.splitlines())
    assert all(c.location == {"sheet": "Q1"} for c in chunks)


def test_chunk_lines_windows_a_single_line_longer_than_size():
    assert chunk_lines("ab\n" + "x" * 10, size=6, overlap=1) == ["ab", "xxxxxx", "xxxxx"]


@pytest.mark.parametrize(
    ("rows", "reason"),
    [
        ([[None, None], ["a", "b"]], "row 1 is empty"),
        ([["name", None], ["a", "b"]], "B1 is empty"),
    ],
)
def test_load_xlsx_raises_with_path_and_sheet_on_bad_header(tmp_path, rows, reason):
    path = tmp_path / "a.xlsx"
    write_xlsx(path, rows)

    with pytest.raises(ValueError, match=f"a.xlsx \\[sheet Q1\\]: .*{reason}"):
        load_xlsx(path)


def test_load_xlsx_raises_on_merged_cells(tmp_path):
    path = tmp_path / "a.xlsx"
    workbook = write_xlsx(path, [["name", "days"], ["An", 12]])
    workbook.active.merge_cells("A2:B2")
    workbook.save(path)

    with pytest.raises(ValueError, match=r"\[sheet Q1\]: merged cells"):
        load_xlsx(path)


def test_load_documents_rejects_xls_and_corrupt_xlsx_naming_the_file(tmp_path):
    (tmp_path / "old.xls").write_bytes(b"x")
    with pytest.raises(LoadError, match="old.xls"):
        load_documents(tmp_path)

    (tmp_path / "old.xls").unlink()
    (tmp_path / "bad.xlsx").write_bytes(b"not a zip")
    with pytest.raises(LoadError, match="bad.xlsx"):
        load_documents(tmp_path)


def test_load_error_leaves_existing_store_untouched(settings, docs_dir):
    ingest(build_chunks(docs_dir, 40, 10), settings, fake_embed, corpus_hash(docs_dir))
    before = open_collection(settings).count()
    (docs_dir / "bad.xlsx").write_bytes(b"not a zip")

    with pytest.raises(LoadError):
        build_chunks(docs_dir, 40, 10)

    assert open_collection(settings).count() == before


def test_load_documents_normalizes_nfd_text_to_nfc(tmp_path):
    (tmp_path / "a.md").write_text(unicodedata.normalize("NFD", "nghỉ phép"), encoding="utf-8")

    assert load_documents(tmp_path)["a.md"][0].text == unicodedata.normalize("NFC", "nghỉ phép")


def test_load_documents_skips_office_lock_files_and_unknown_suffixes(tmp_path):
    (tmp_path / "~$a.docx").write_bytes(b"lock")
    (tmp_path / "a.csv").write_text("x", encoding="utf-8")
    (tmp_path / "a.md").write_text("hello", encoding="utf-8")

    assert list(load_documents(tmp_path)) == ["a.md"]


def test_load_documents_rejects_non_utf8_text_naming_the_file(tmp_path):
    (tmp_path / "a.md").write_bytes(b"\xff\xfe\x00bad")

    with pytest.raises(LoadError, match="a.md"):
        load_documents(tmp_path)


def test_corpus_hash_changes_with_one_byte_and_ignores_unsupported_files(tmp_path):
    (tmp_path / "a.md").write_text("one", encoding="utf-8")
    first = corpus_hash(tmp_path)

    (tmp_path / "ignored.csv").write_text("x", encoding="utf-8")
    assert corpus_hash(tmp_path) == first

    (tmp_path / "a.md").write_text("onf", encoding="utf-8")
    assert corpus_hash(tmp_path) != first
