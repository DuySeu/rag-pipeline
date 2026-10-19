from __future__ import annotations

import logging
import zipfile
from collections.abc import Callable, Sequence
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from docx import Document
from docx.table import Table
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.shapes.base import BaseShape
from pypdf import PdfReader

logger = logging.getLogger(__name__)

NOTES_LABEL = "Ghi chú:"


# A unit of text that chunking never crosses: a page, a slide, a sheet or a whole file
@dataclass(frozen=True)
class Section:
    text: str
    location: dict[str, str | int] = field(default_factory=dict)
    line_wise: bool = False  # chunk by packing whole lines instead of a character window


# A file that cannot be read or is outside what its loader supports
class LoadError(ValueError):
    pass


# Read a Markdown or plain-text file as one section
def load_text(path: Path) -> list[Section]:
    return [Section(path.read_text(encoding="utf-8"))]


# Read paragraphs and tables of a Word file in body order as one section
def load_docx(path: Path) -> list[Section]:
    lines: list[str] = []
    for block in Document(str(path)).iter_inner_content():
        if isinstance(block, Table):
            lines.extend(" | ".join(cell.text for cell in row.cells) for row in block.rows)
        else:
            lines.append(block.text)
    return [Section("\n".join(line for line in lines if line.strip()))]


# Collect the text of shapes in order, descending into groups and tables
def slide_lines(shapes: Sequence[BaseShape], skip_id: int | None) -> list[str]:
    lines: list[str] = []
    for shape in shapes:
        if shape.shape_id == skip_id:
            continue
        if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
            lines.extend(slide_lines(list(shape.shapes), skip_id))
        elif shape.has_table:
            lines.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
        elif shape.has_text_frame:
            lines.append(shape.text_frame.text)
    return lines


# Read a PowerPoint file as one section per slide: title, shapes, then labelled speaker notes
def load_pptx(path: Path) -> list[Section]:
    sections: list[Section] = []
    for number, slide in enumerate(Presentation(str(path)).slides, start=1):
        title = slide.shapes.title
        lines = [title.text_frame.text] if title is not None else []
        lines += slide_lines(list(slide.shapes), title.shape_id if title is not None else None)

        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        if notes.strip():
            lines += [NOTES_LABEL, notes]
        sections.append(
            Section("\n".join(line for line in lines if line.strip()), {"slide": number})
        )
    return sections


# Read a PDF as one section per page, warning about pages without a text layer
def load_pdf(path: Path) -> list[Section]:
    sections: list[Section] = []
    for number, page in enumerate(PdfReader(path).pages, start=1):
        text = page.extract_text() or ""
        if not text.strip():
            logger.warning(
                "%s page %d has no text layer, skipped (OCR is not supported)", path, number
            )
        sections.append(Section(text, {"page": number}))
    return sections


# Read every visible sheet of a basic Excel table (headers in row 1, one table from A1) as rows
def load_xlsx(path: Path) -> list[Section]:
    try:
        workbook = openpyxl.load_workbook(path, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise LoadError(f"{path}: corrupt or password-protected file ({exc})") from exc

    sections: list[Section] = []
    for sheet in workbook.worksheets:
        if sheet.sheet_state != "visible":
            logger.info("%s [sheet %s]: hidden sheet skipped", path, sheet.title)
            continue

        rows = list(sheet.iter_rows(values_only=True))
        if not any(cell not in (None, "") for row in rows for cell in row):
            continue

        # Fail on layouts the row-per-record rendering would silently get wrong
        where = f"{path} [sheet {sheet.title}]"
        if sheet.merged_cells.ranges:
            raise LoadError(f"{where}: merged cells are not supported")
        header = ["" if cell is None else str(cell).strip() for cell in rows[0]]
        if not any(header):
            raise LoadError(f"{where}: row 1 is empty, headers must be in row 1")
        data = [row for row in rows[1:] if any(cell not in (None, "") for cell in row)]
        for column, name in enumerate(header):
            if not name and any(row[column] not in (None, "") for row in data):
                raise LoadError(f"{where}: header cell {get_column_letter(column + 1)}1 is empty")

        # One line per record, repeating the headers so each line stands alone
        lines = [
            f"Sheet {sheet.title}: "
            + "; ".join(
                f"{name}: {' '.join(str(cell).split())}"
                for name, cell in zip(header, row, strict=False)
                if cell not in (None, "")
            )
            for row in data
        ]
        sections.append(Section("\n".join(lines), {"sheet": sheet.title}, line_wise=True))
    return sections


# Reject the legacy Excel format loudly instead of skipping it
def reject_xls(path: Path) -> list[Section]:
    raise LoadError(f"{path}: .xls is not supported, convert it to .xlsx")


LOADERS: dict[str, Callable[[Path], list[Section]]] = {
    ".md": load_text,
    ".txt": load_text,
    ".docx": load_docx,
    ".xlsx": load_xlsx,
    ".xls": reject_xls,
    ".pptx": load_pptx,
    ".pdf": load_pdf,
}
